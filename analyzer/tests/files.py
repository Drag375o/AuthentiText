"""Builders for real test documents, generated in memory."""
import io
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile

PARAGRAPH_1 = "My grandmother kept her recipes on the backs of electricity bills."
PARAGRAPH_2 = "Half of them are unreadable now, stained with turmeric."


def upload(name: str, data: bytes, content_type: str = "application/octet-stream") -> SimpleUploadedFile:
    return SimpleUploadedFile(name, data, content_type=content_type)


def make_pdf(paragraphs=(PARAGRAPH_1, PARAGRAPH_2), password: str | None = None, pages: int = 1, text=True) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page()
        if text:
            y = 72
            for para in paragraphs:
                rect = pymupdf.Rect(72, y, 300, y + 120)   # narrow box forces line wrapping
                page.insert_textbox(rect, para, fontsize=11)
                y += 140
        else:
            page.draw_rect(pymupdf.Rect(72, 72, 300, 300), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    kwargs = {}
    if password:
        kwargs = {"encryption": pymupdf.PDF_ENCRYPT_AES_256, "user_pw": password, "owner_pw": password}
    data = doc.tobytes(**kwargs)
    doc.close()
    return data


def make_docx(paragraphs=(PARAGRAPH_1, PARAGRAPH_2), table: bool = False) -> bytes:
    import docx
    document = docx.Document()
    for para in paragraphs:
        document.add_paragraph(para)
    if table:
        document.add_table(rows=2, cols=2).cell(0, 0).text = "Table text"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def add_to_zip(data: bytes, name: str, content: bytes) -> bytes:
    buffer = io.BytesIO(data)
    with zipfile.ZipFile(buffer, "a") as archive:
        archive.writestr(name, content)
    return buffer.getvalue()


def make_zip_bomb(uncompressed_mb: int = 60) -> bytes:
    """A .docx-shaped zip whose content expands far beyond its size on disk."""
    data = make_docx()
    buffer = io.BytesIO(data)
    with zipfile.ZipFile(buffer, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/media/padding.bin", b"\0" * (uncompressed_mb * 1024 * 1024))
    return buffer.getvalue()


EXE_BYTES = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00" + b"\x00" * 64
